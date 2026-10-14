#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
stock-master 自检
=================
逐个调用数据层 / 分析层 / 筛选层的公开接口，报告可用性。
用法：
    python3 scripts/_selftest.py            # 全量
    python3 scripts/_selftest.py --quick    # 跳过慢项
退出码：0 全通过 / 1 有失败项
"""
import sys
import time
import traceback
from datetime import datetime

sys.path.insert(0, "/var/minis/skills/stock-master")

RESULTS = []
QUICK = "--quick" in sys.argv


def check(name, fn, quick=False):
    """跑一个检查项。quick=True 的项在 --quick 模式下跳过。"""
    if QUICK and quick:
        RESULTS.append((name, "SKIP", ""))
        print(f"  ⊘ {name}（跳过）")
        return None
    t0 = time.time()
    try:
        val = fn()
        dt = time.time() - t0
        ok = val is not False and val is not None
        if val == [] or val == {}:
            ok = False
        summary = _summary(val)
        RESULTS.append((name, "OK" if ok else "EMPTY", summary))
        print(f"  {'✅' if ok else '⚠️ '} {name}  [{dt:.1f}s]  {summary}")
        return val
    except Exception as e:                                       # noqa: BLE001
        dt = time.time() - t0
        box = f"{type(e).__name__}: {e}"
        RESULTS.append((name, "FAIL", box))
        print(f"  ❌ {name}  [{dt:.1f}s]  {box}")
        if "--verbose" in sys.argv:
            traceback.print_exc()
        return None


def _summary(v):
    if isinstance(v, list):
        first = v[0] if v else None
        if isinstance(first, dict):
            keys = ",".join(list(first.keys())[:4])
            return f"{len(v)} 条 | {keys}"
        return f"{len(v)} 条 | {str(first)[:40]}"
    if isinstance(v, dict):
        return f"{len(v)} 键 | {','.join(list(v.keys())[:5])}"
    if isinstance(v, (str, int, float)):
        return str(v)[:60]
    return type(v).__name__


def section(title):
    print(f"\n{'=' * 62}\n{title}\n{'=' * 62}")


# ============================================================
# 1. 交易日 / 日历
# ============================================================
section("1. 交易日历")

from webtool.stock.skill.scripts import short_tool as ST                               # noqa: E402

check("short_tool.calendar", lambda: ST.calendar())

# ============================================================
# 2. 行情（腾讯直连）
# ============================================================
section("2. 行情 · 腾讯直连")

q = check("short_tool.quote 个股", lambda: ST.quote(["600519", "600121"]))
if q:
    def _assert_limit():
        bad = []
        for r in q:
            if not r.get("ok"):
                bad.append(r["input"])
                continue
            exp_zt = round(r["昨收"] * 1.1, 2)
            exp_dt = round(r["昨收"] * 0.9, 2)
            if abs(r["涨停价"] - exp_zt) > 0.03 or abs(r["跌停价"] - exp_dt) > 0.03:
                bad.append(f"{r['input']}(涨停{r['涨停价']}≠{exp_zt})")
            if not (0 < r["现价"] < 100000):
                bad.append(f"{r['input']}(现价异常{r['现价']})")
        assert not bad, f"字段校验失败：{bad}"
        return "涨停/跌停/现价 字段断言通过"
    check("  字段断言（涨停=昨收×1.1）", _assert_limit)

check("short_tool.quote 指数", lambda: ST.quote(
    ["上证指数", "深证成指", "创业板指", "科创50", "沪深300", "中证500"]))
check("short_tool.quote 异常码容错", lambda: [
    r for r in ST.quote(["000000", "999999"]) if r])

# ============================================================
# 3. 涨停池
# ============================================================
section("3. 涨停池 · 东财 push2ex")

pool = check("short_tool.pool（全量）", lambda: ST.pool("last", "all"))
if pool:
    def _pool_check():
        assert pool["涨停总数"] > 0, "涨停总数为 0"
        r = pool["rows"][0]
        for k in ("code", "name", "连板", "首封", "封单亿", "可买"):
            assert k in r, f"缺字段 {k}"
        assert isinstance(r["连板"], int) and r["连板"] >= 1, "连板数异常"
        return (f"涨停 {pool['涨停总数']} / 可买 {pool['主板可买数']} "
                f"({pool['主板可买占比']}) 梯队 {pool['连板梯队']}")
    check("  涨停池字段完整性", _pool_check)
check("short_tool.pool（只看可买）", lambda: ST.pool("last", "buy"))

# ============================================================
# 4. 可买入性预检
# ============================================================
section("4. 可买入性预检 / 板块过滤")

check("short_tool.check", lambda: ST.check(["600886", "600621", "601225"]))
check("short_tool.filter 主板", lambda: ST.filter(
    ["600121", "002068", "000001", "601398"]))
check("short_tool.filter 不可买", lambda: ST.filter(
    ["300913", "301001", "688559", "920157", "900901", "200011"]))


def _filter_neg():
    res = {r["code"]: r["通过"] for r in ST.filter(
        ["300913", "688559", "920157", "900901", "200011", "430047"])}
    wrong = [c for c, ok in res.items() if ok]
    assert not wrong, f"以下应为不可买却通过：{wrong}"
    return "6 类不可买板块全部正确拦截"


check("  不可买板块拦截（负例）", _filter_neg)


def _st_blocked():
    """ST / 退市整理必须一票否决，且不能被『无涨跌停』分支漏掉"""
    rows = {r["input"]: r for r in ST.check(["000004", "600666"])}
    st = rows.get("000004", {})
    assert "❌" in (st.get("权限") or ""), f"退市股未被判不可买：{st.get('权限')}"
    assert st.get("结论", "").startswith("❌"), f"结论未否决：{st.get('结论')}"
    return f"000004 {st.get('name')} → {st.get('权限')} / {st.get('结论')}"


check("  退市/ST 一票否决（负例）", _st_blocked)


def _pool_carries_name():
    """涨停池必须把名称传给可买性判定，否则 ST 只能靠代码段识别"""
    r = ST.pool("last", "all")
    rows = r.get("rows") or []
    assert rows, "涨停池为空"
    noname = [x["code"] for x in rows if not x.get("name")]
    assert not noname, f"涨停池缺名称：{noname[:5]}"
    return f"{r['date']} 涨停 {r['涨停总数']} 只，主板可买 {r['主板可买数']}"


check("  涨停池名称透传", _pool_carries_name)


def _pool_date_guard():
    """东财忽略传入日期, 回的是"最新一个有数据的交易日"。请求非交易日时必须
    标出实际数据日并告警, 不能假装那就是请求日的数据。

    断言不写死具体日期 (写死只在当天成立): 只要求回退到的是一个真实交易日,
    且不等于请求日。
    """
    r = ST.pool("2026-10-11", "all")        # 周末, 非交易日
    assert r.get("日期回退"), "非交易日请求未告警"
    assert r["date"] != "20261011", f"东财居然认了非交易日: {r['date']}"
    d = r["date"]
    assert datetime.strptime(d, "%Y%m%d").date().weekday() < 5, f"回退到的不是交易日: {d}"
    # 反过来: 请求这个"最新数据日"本身 → 不该再报回退
    r2 = ST.pool(f"{d[:4]}-{d[4:6]}-{d[6:]}", "all")
    assert not r2.get("日期回退"), f"请求最新数据日不该报回退: {r2.get('date')}"
    return r["日期回退"]


check("  涨停池日期回退告警", _pool_date_guard)
check("short_tool.header screener", lambda: ST.header("pre", "screener", None))
check("short_tool.header limitup", lambda: ST.header("after", "limitup", None))

# ============================================================
# 5. clients 数据层
# ============================================================
section("5. clients 数据层")

from webtool.stock.skill.scripts import clients as C                                   # noqa: E402

check("clients.quote", lambda: C.quote(["600519"]))
check("clients.limit_up_pool", lambda: C.limit_up_pool("last"))
check("clients.prev_trading_day", lambda: C.prev_trading_day())
check("clients.is_trading_day", lambda: "今天是交易日" if C.is_trading_day(
    datetime.now()) else "今天非交易日", quick=True)
check("clients.news", lambda: C.news("600519"), quick=True)
check("clients.announcements（巨潮）", lambda: C.announcements("600519", 30))
check("clients.code_ok", lambda: C.code_ok("600121"))
check("clients.web_search", lambda: C.web_search("上证指数 今日", 1), quick=True)

# ============================================================
# 6. feeds 层直连 (生产路径: analysis/、screening/ 都走 feeds)
# ============================================================
section("6. feeds 直连")

from webtool.stock.skill.scripts import feeds as F                                     # noqa: E402

check("feeds.quote 单只", lambda: F.quote("600519"))
check("feeds.quote 批量", lambda: F.quote(["600519", "600121"]))
check("feeds.kline_list", lambda: F.kline_list("600519", count=30))
check("feeds.rank_list 涨幅榜", lambda: F.rank_list("涨幅榜", 50))
check("feeds.rank_list 成交额榜", lambda: F.rank_list("成交额榜", 30))
check("feeds.announcements 平安银行", lambda: F.announcements("000001", 30))
check("feeds.announcements 中信证券", lambda: F.announcements("600030", 30), quick=True)
check("feeds.industry_rows", lambda: F.industry_rows(10), quick=True)
check("feeds.industry_kline", lambda: F.industry_kline("881101"), quick=True)
check("feeds.company_intro", lambda: F.company_intro("600519"), quick=True)


def _kline_span():
    """K 线跨度必须按 count 反推自然日：否则请求 250 条只回 58 条，
    MA60/年线等长周期指标会静默算错（历史上正是这个坑）"""
    n = len(F.kline_list("600519", count=250)['items'])
    assert n >= 250, f"请求 250 条只回 {n} 条，跨度不足"
    return f"请求 250 → 实得 {n} 条"


check("feeds.kline 跨度充足", _kline_span, quick=True)

# ============================================================
# 7. 分析层
# ============================================================
section("7. 分析层")

from webtool.stock.skill.scripts.analysis import tech_indicators as TI                  # noqa: E402
from webtool.stock.skill.scripts.analysis import single_analysis as SA2                 # noqa: E402

def _score_range():
    r = SA2.analyze("600519")
    s = r["score"]
    assert s is not None and 0 <= s <= 100, f"评分越界 {s}"
    bd = r["score_breakdown"]
    tot = sum(bd.values())
    assert abs(tot - s) < 0.15, f"分解和 {tot} ≠ 总分 {s}"
    assert "技术面(60)" in bd and "消息面(30)" in bd and "宏观(10)" in bd, f"分解缺项 {bd}"
    return f"评分 {s} = {bd}"


check("single_analysis 100 分制断言", _score_range, quick=True)
check("tech_indicators.score_short", lambda: TI.score_short(
    TI.build_df("600519") if hasattr(TI, "build_df") else None), quick=True)

# ============================================================
# 8. 筛选层
# ============================================================
section("8. 筛选层")

from webtool.stock.skill.scripts.screening import risk_check as RC                      # noqa: E402


def _risk_positive():
    """已知含预亏公告的标的应被命中（正例）"""
    hits = []
    for code in ("600121", "600207"):
        r = RC.scan(code)
        if r["hit"]:
            hits.append((code, r["hit"][0]["keyword"]))
    assert hits, "已知风险标的均未命中（关键词表失效？）"
    return f"命中 {hits}"


def _risk_negative():
    """干净标的应无命中（负例）"""
    r = RC.scan("600519")
    assert not r["hit"], f"茅台被误判：{r['hit']}"
    return "无误判"


def _risk_fundamentals():
    """基本面红线：郑州煤电资产负债率 84% + 净利同比为负，必须命中"""
    r = RC.scan("600121")
    kw = " ".join(h["keyword"] for h in r["hit"] if h["type"] == "基本面")
    assert "资产负债率" in kw, f"基本面红线未生效：{r['hit']}"
    return kw


check("risk_check 正例", _risk_positive)
check("risk_check 基本面红线", _risk_fundamentals)
check("risk_check 负例", _risk_negative)


def _risk_noise():
    """申报材料类公告不得误判为风险（噪声豁免）"""
    good = "关于最近五年被证券监管部门和交易所采取监管措施或处罚情况的公告"
    assert any(p.search(good) for p in RC.NOISE_PATTERNS), "申报附件未被豁免"
    bad = "关于收到深圳证券交易所信息披露监管问询函的公告"
    assert not any(p.search(bad) for p in RC.NOISE_PATTERNS), "真问询函被误豁免"
    return "申报附件豁免 / 真问询函保留"


check("risk_check 噪声豁免", _risk_noise, quick=True)

# ============================================================
# 9. 长中线池 / 典型标的池可买性
# ============================================================
section("9. 长中线池 / 典型标的池")

from webtool.stock.skill.scripts.screening import mid_pool as MP                        # noqa: E402
from webtool.stock.skill.scripts.screening import long_pool as LP                       # noqa: E402
from webtool.stock.skill.scripts.screening import pool_selector as PS                   # noqa: E402
from webtool.stock.skill.scripts.screening import short_pool as SP_short                # noqa: E402
from webtool.stock.skill.scripts.export import short_screen_excel as EX2               # noqa: E402
from webtool.stock.skill.scripts._common import is_tradable                             # noqa: E402


def _typical_all_tradable():
    """典型标的池不允许含不可买板块（300/301/688/920/ST）"""
    tables = {}
    tables.update({f"中线/{k}": v for k, v in MP.suggest_typical_stocks().items()})
    tables.update({f"长线/{k}": v for k, v in LP.suggest_typical_stocks().items()})
    bad = [c for lst in tables.values() for c in lst
           if not is_tradable(c.split()[0])]
    assert not bad, f"典型池含不可买标的：{bad}"
    return f"{len(tables)} 组 / {sum(len(v) for v in tables.values())} 只，全部主板可买"


check("典型标的池可买性", _typical_all_tradable, quick=True)


def _pool_route_long():
    """pool_selector → 长线池：相同入参两种写法都要能吃"""
    r1 = PS.select("long", codes=[{"code": "600519", "name": "贵州茅台"}])
    assert r1["passed"] or r1["failed"], "长线池返回空"
    return f"长线 通过 {len(r1['passed'])} / 未过 {len(r1['failed'])}"


check("pool_selector 长线池", _pool_route_long)


def _pool_route_mid():
    """pool_selector → 中线池：裸码列表（旧接口会 AttributeError）"""
    r = PS.select("mid", codes=["600519"])
    assert r["passed"] or r["failed"], "中线池返回空"
    return f"中线 通过 {len(r['passed'])} / 未过 {len(r['failed'])}"


check("pool_selector 中线池（裸码入参）", _pool_route_mid, quick=True)

def _pool_portfolio_layer():
    """组合层 + Excel 一键导出（共用一次筛选，避免重复跑 80s 的短池）"""
    r = SP_short.run(top_n_movers=30, top_n_sectors=6, min_score=60)
    p = r["passed"][0]
    assert 0.0 <= p["position_60d"] <= 1.0, f"位置分位越界：{p['position_60d']}"
    assert p["position_tag"].startswith(("高位", "中位", "低位")), p["position_tag"]
    assert p["position_window"] > 0
    assert "industries" in r and "portfolio_notes" in r

    path = EX2.build_from_screen(r, "/tmp/_selftest_short.xlsx")
    import openpyxl
    wb = openpyxl.load_workbook(path)
    assert "板块行业趋势" in wb.sheetnames, wb.sheetnames
    first = str(wb["板块行业趋势"].cell(row=2, column=1).value or "")
    assert "申万一级覆盖" in first, f"板块总评未整行写入：{first[:40]}"

    return (f"通过 {r['passed_count']} / 行业 {r['industry_dispersion']} 个 / "
            f"高位 {len(r['high_position_codes'])} 只 / Excel {len(wb.sheetnames)} 表")


check("short_pool 组合层 + Excel 一键导出", _pool_portfolio_layer)




# ============================================================
# 10. 情绪 / 可视化 / 研报
# ============================================================
section("10. 情绪 / 可视化 / 研报")

from webtool.stock.skill.scripts.sentiment import news_scorer as NS                     # noqa: E402
from webtool.stock.skill.scripts.visualization import logic_renderer as LR              # noqa: E402
from webtool.stock.skill.scripts.research import report_writer as RW                    # noqa: E402


def _news_scorer():
    pos = NS.score(["公司净利润大增，订单超预期"])["score"]
    neg = NS.score(["业绩预亏，计提减值"])["score"]
    assert pos > 0 > neg, f"情绪方向错误 pos={pos} neg={neg}"
    return f"正 {pos} / 负 {neg}，方向正确"


check("news_scorer 方向性", _news_scorer, quick=True)


def _renderer():
    md = LR.decision_tree()
    assert "市值 50-300" not in md, "决策树仍含已废弃的 300 亿上限"
    assert "候选池" in md and "```mermaid" in md
    # mindmap 的 ((...)) 内不能有换行/半角括号（mermaid 会解析失败）
    mm = LR.investment_logic("600519", "贵州茅台", "高股息(稳健)",
                             ["提价预期(直营)"], ["消费税调整"], ["三季报"])
    root = [ln for ln in mm.split("\n") if "root((" in ln][0]
    assert "\\n" not in root and "(" not in root.split("root((")[1][:-2]
    # 产业链图方向必须 上游→中游→下游
    ic = LR.industry_chain("白酒", ["高粱"], ["酿造"], ["商超"])
    assert "U[上游" in ic and "-->" in ic and ic.count("subgraph") == 1
    return "决策树 / mindmap / 产业链 三图结构合规"


check("logic_renderer 三图结构", _renderer, quick=True)


def _report():
    r = SA2.analyze("600519")
    txt = RW.single_stock_report(r)
    assert "/100" in txt and "综合评分" in txt, "研报未标明 100 分制量纲"
    # 短选报告：industry 现为字符串（申万一级），不能再按列表取
    txt2 = RW.short_screen_report({
        "pool_size": 84, "passed_count": 1, "min_score_threshold": 60, "failed": [],
        "passed": [{"code": "600640", "name": "国脉文化", "score": 94,
                    "float_market_capital_yi": 60.0, "industry": "传媒",
                    "technical": {"current_price": 12.3, "entry_breakout": 12.5,
                                  "entry_pullback": 11.8, "hard_stop": 11.4}}]})
    assert "国脉文化" in txt2 and "传媒" in txt2, "短选报告渲染失败"
    return f"单股 {len(txt)} / 短选 {len(txt2)} 字符"


check("report_writer 量纲", _report, quick=True)

# ============================================================
# 汇总
# ============================================================
section("汇总")
ok = sum(1 for _, s, _ in RESULTS if s == "OK")
empty = sum(1 for _, s, _ in RESULTS if s == "EMPTY")
fail = sum(1 for _, s, _ in RESULTS if s == "FAIL")
skip = sum(1 for _, s, _ in RESULTS if s == "SKIP")
print(f"  ✅ OK {ok}   ⚠️ 空 {empty}   ❌ 失败 {fail}   ⊘ 跳过 {skip}")
if fail or empty:
    print("\n  需修复：")
    for n, s, m in RESULTS:
        if s in ("FAIL", "EMPTY"):
            print(f"    [{s}] {n} — {m}")
sys.exit(1 if (fail or empty) else 0)
